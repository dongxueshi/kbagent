import io
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ImportFailure(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".md", ".txt"}:
            content = data.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ImportFailure("请先解除 PDF 加密再上传")
            if len(reader.pages) > 500:
                raise ImportFailure("PDF 最多支持 500 页")
            content = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(z.file_size for z in archive.infolist()) > 50 * 1024 * 1024:
                    raise ImportFailure("DOCX 解压后超过 50 MB 限制")
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            parts += [
                " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
            ]
            content = "\n\n".join(parts)
        else:
            raise ImportFailure("支持的格式为 .md、.txt、.pdf、.docx")
    except ImportFailure:
        raise
    except Exception as exc:
        raise ImportFailure("无法读取文件，请检查文件格式；文本文件应使用 UTF-8 编码") from exc
    content = content.replace("\x00", "").strip()
    if not content:
        raise ImportFailure("文件中没有可检索文本；扫描版 PDF 请先进行 OCR")
    if len(content) > 200_000:
        raise ImportFailure("文档超过 200,000 字符，请拆分后上传")
    return Path(filename).stem[:200] or "导入文档", content
