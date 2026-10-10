import io
import json

import pytest
from docx import Document

from kbagent.main import create_app


def events(response):
    return [
        json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")
    ]


def test_health_empty_state_and_secrets(client):
    assert client.get("/api/health").json() == {
        "status": "ok",
        "database": "ok",
        "object_storage": "ok",
    }
    assert client.get("/api/spaces").json() == []
    assert client.get("/").status_code == 200
    assert client.get("/static/app.js").status_code == 200
    data = client.get("/api/settings").json()
    assert data["mode"] == "retrieval"
    assert set(data) == {"mode", "model", "sdk", "database", "object_storage", "max_upload_bytes"}
    assert data["database"] == "mysql" and data["object_storage"] == "minio"


def test_chinese_search_update_delete_and_isolation(client, base):
    doc = client.post(
        base + "/documents", json={"title": "部署指南", "content": "安装开发环境需要运行 uv sync。"}
    ).json()
    hits = client.get(base + "/search", params={"q": "开发环境"}).json()
    assert hits and hits[0]["document_id"] == doc["id"]
    other = client.post("/api/spaces", json={"name": "其他库"}).json()
    other_base = f"/api/spaces/{other['id']}"
    assert client.get(other_base + "/search", params={"q": "开发环境"}).json() == []
    assert client.get(other_base + "/documents/" + doc["id"]).status_code == 404
    assert client.delete(other_base + "/documents/" + doc["id"]).status_code == 404
    assert (
        client.put(
            base + "/documents/" + doc["id"],
            json={"title": "报销指南", "content": "交通报销需要发票。"},
        ).status_code
        == 200
    )
    assert client.get(base + "/search", params={"q": "开发环境"}).json() == []
    assert client.get(base + "/search", params={"q": "交通报销"}).json()
    exported = client.get(base + "/documents/" + doc["id"] + "/export")
    assert "交通报销" in exported.text
    assert client.delete(base + "/documents/" + doc["id"]).status_code == 204
    assert client.get(base + "/search", params={"q": "报销"}).json() == []


def test_upload_text_docx_and_rejections(client, base):
    uploaded = client.post(
        base + "/import",
        files={"file": ("手册.md", "# 文档导入\n支持知识库检索。".encode(), "text/markdown")},
    )
    assert uploaded.status_code == 201
    assert uploaded.json()["title"] == "手册"
    word = Document()
    word.add_paragraph("Word 文件内容")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "预算"
    table.cell(0, 1).text = "5000元"
    buf = io.BytesIO()
    word.save(buf)
    response = client.post(base + "/import", files={"file": ("预算.docx", buf.getvalue())})
    assert response.status_code == 201
    assert "预算 | 5000元" in response.json()["content"]
    assert (
        client.get(base + f"/documents/{response.json()['id']}/original").content == buf.getvalue()
    )
    for name, content in [
        ("invalid.pdf", b"broken"),
        ("empty.txt", b""),
        ("bad.txt", b"\xff"),
        ("script.exe", b"abc"),
    ]:
        assert client.post(base + "/import", files={"file": (name, content)}).status_code == 422


def test_upload_size_limit(client, base):
    client.app.state.settings.max_upload_bytes = 8
    assert (
        client.post(base + "/import", files={"file": ("large.txt", b"123456789")}).status_code
        == 413
    )


def test_pdf_text_extraction(client, base):
    # A minimal real PDF with extractable text and a valid cross-reference table.
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    content = b"BT /F1 12 Tf 72 720 Td (Knowledge import guide) Tj ET"
    objects.append(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
    pdf = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(pdf)
    pdf += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    pdf += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets[1:])
    pdf += f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    response = client.post(base + "/import", files={"file": ("guide.pdf", pdf)})
    assert response.status_code == 201
    assert "Knowledge import guide" in response.json()["content"]
    assert client.get(base + f"/documents/{response.json()['id']}/original").content == pdf


def test_retrieval_chat_persistence_and_cascade(client, base, space):
    doc = client.post(
        base + "/documents", json={"title": "模型配置", "content": "模型需要支持工具调用。"}
    ).json()
    conversation = client.post(base + "/conversations").json()
    path = base + "/conversations/" + conversation["id"]
    response = client.post(path + "/chat", json={"question": "模型配置"})
    result = events(response)[-1]
    assert result["type"] == "answer" and result["mode"] == "retrieval"
    assert result["sources"][0]["document_id"] == doc["id"]
    assert "并非 AI 生成答案" in result["content"]
    history = client.get(path).json()
    assert [m["role"] for m in history["messages"]] == ["user", "assistant"]
    assert history["title"] == "模型配置"
    other = client.post("/api/spaces", json={"name": "隔离库"}).json()
    assert (
        client.get(f"/api/spaces/{other['id']}/conversations/{conversation['id']}").status_code
        == 404
    )
    result = events(client.post(path + "/chat", json={"question": "火星航天器"}))[-1]
    assert result["sources"] == []
    assert "没有找到" in result["content"]
    assert client.delete(base).status_code == 204
    assert client.get(path).status_code == 404
    with client.app.state.store.connection() as db:
        for table in ["documents", "chunks", "chunk_terms", "conversations", "messages"]:
            assert db.execute(f"SELECT count(*) FROM {table}").fetchone()[0] == 0


def test_restart_retains_data(client, base):
    from fastapi.testclient import TestClient

    client.post(base + "/documents", json={"title": "持久化", "content": "重启不会丢失资料"})
    uploaded = client.post(
        base + "/import", files={"file": ("restart.txt", b"persistent original")}
    ).json()
    with TestClient(create_app(client.app.state.settings)) as second:
        assert len(second.get(base + "/documents").json()) == 2
        assert (
            second.get(base + f"/documents/{uploaded['id']}/original").content
            == b"persistent original"
        )
        assert second.get(base + "/search", params={"q": "重启"}).json()


@pytest.mark.parametrize(
    "body",
    [
        {"title": "", "content": "x"},
        {"title": " ", "content": "x"},
        {"title": "x", "content": " "},
        {"title": "x" * 201, "content": "x"},
    ],
)
def test_invalid_document(client, base, body):
    assert client.post(base + "/documents", json=body).status_code == 422


def test_search_query_escaping(client, base):
    client.post(base + "/documents", json={"title": "FTS", "content": "报价单价格"})
    for query in ['" OR *', '"报价" AND (价格)', ':"; DROP TABLE documents; --']:
        assert client.get(base + "/search", params={"q": query}).status_code == 200
    assert len(client.get(base + "/documents").json()) == 1
